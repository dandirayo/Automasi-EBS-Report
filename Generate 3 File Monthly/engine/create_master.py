import os
from docx import Document
from pptx import Presentation

def replace_text_in_docx(doc_path, output_path, mapping):
    doc = Document(doc_path)
    for p in doc.paragraphs:
        for key, val in mapping.items():
            if key in p.text:
                for run in p.runs:
                    if key in run.text:
                        run.text = run.text.replace(key, val)
                if key in p.text:
                    p.text = p.text.replace(key, val)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for key, val in mapping.items():
                        if key in p.text:
                            for run in p.runs:
                                if key in run.text:
                                    run.text = run.text.replace(key, val)
                            if key in p.text:
                                p.text = p.text.replace(key, val)
    doc.save(output_path)

def replace_text_in_pptx(ppt_path, output_path, mapping):
    prs = Presentation(ppt_path)
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                for key, val in mapping.items():
                    if key in shape.text:
                        for paragraph in shape.text_frame.paragraphs:
                            for run in paragraph.runs:
                                if key in run.text:
                                    run.text = run.text.replace(key, val)
                        if key in shape.text:
                            shape.text = shape.text.replace(key, val)
            if shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        for key, val in mapping.items():
                            if key in cell.text:
                                for paragraph in cell.text_frame.paragraphs:
                                    for run in paragraph.runs:
                                        if key in run.text:
                                            run.text = run.text.replace(key, val)
                                if key in cell.text:
                                    cell.text = cell.text.replace(key, val)
    prs.save(output_path)

def main():
    template_dir = r"C:\Users\901191\Music\Monthly_Report Downloaded-Edited\OneDrive\Juli"
    output_dir = r"C:\Users\901191\Music\Automasi EFS\Generate 3 File Monthly\Template_Master"
    
    template_docx = os.path.join(template_dir, "Draft Notin Penyampaian Monthly Report CBS Juli- Enterprise Financial System (EFS) Salin.docx")
    template_pptx = os.path.join(template_dir, "202607 - EFS Monthly Report - July 2026.pptx")
    
    output_docx = os.path.join(output_dir, "Master_Template_Notin.docx")
    output_pptx = os.path.join(output_dir, "Master_Template_Presentation.pptx")
    
    mapping = {
        "Juli 2026": "{{BULAN_TAHUN}}",
        "July 2026": "{{MONTH_YEAR}}",
        "JULY 2026": "{{MONTH_YEAR_UPPER}}",
        "Juli": "{{BULAN}}",
        "July": "{{MONTH}}",
        "JULY": "{{MONTH_UPPER}}",
        "Juni": "{{BULAN_LALU}}",
        "June": "{{MONTH_PREV}}",
        "202607": "{{YYYYMM}}",
        "99,51%": "{{XLA_SUCCESS_PCT}}",
        "99,36%": "{{XLA_SUCCESS_PCT}}",
        "70.091.415": "{{XLA_PROCESSED}}",
        "70,5 juta": "{{XLA_TOTAL_JUTA}}",
        "1.930.602": "{{XLA_UNPROCESSED_PREV}}",
        "449.109": "{{XLA_UNPROCESSED}}",
        "137.003": "{{XLA_ERROR}}",
        "5 / 8": "{{INFRA_CRITICAL}} / 8",
        "4 Critical": "{{INFRA_CRITICAL}} Critical"
    }
    
    replace_text_in_docx(template_docx, output_docx, mapping)
    replace_text_in_pptx(template_pptx, output_pptx, mapping)
    print("Master Templates Created!")

if __name__ == "__main__":
    main()
